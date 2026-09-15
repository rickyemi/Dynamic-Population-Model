"""Difference tab rebuild on a synthetic Dual Use workbook.

The workbook the self test builds carries no Difference tab, which is the case a new product
presents: there is nothing to compare against, so only the structural checks can run. That path has
to work, and it is the one a committed fixture with a Difference tab would never exercise.
"""
from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

from dpm import dual_use_optimizer as du


def test_self_test(workdir):
    du.self_test(workdir / "st")


def test_output_has_exactly_three_tabs(workdir):
    du.build_self_test_workbook(workdir / "in")
    result = du.run_optimizer(data_dir=workdir / "in", out_dir=workdir / "out",
                              validate=False, verbose=False)
    tabs = load_workbook(result["output_file"]).sheetnames
    assert tabs == ["Detailed Results", "MasterModels TP-ID", du.OUTPUT_DIFF_SHEET]


def test_requirement_1_rejects_a_misnamed_file(workdir):
    folder = workdir / "in"
    path = du.build_self_test_workbook(folder)
    path.rename(folder / "DualUse_ProductS.xlsx")
    with pytest.raises(FileNotFoundError, match="Requirement 1"):
        du.resolve_input(folder)


def test_requirement_3_rejects_a_model_group_without_initials():
    with pytest.raises(ValueError, match="Requirement 3"):
        du.split_model_group("ProductS")
    assert du.split_model_group("Product_S_YO") == ("Product_S", "YO")


def test_tab_names_match_loosely():
    """'Raw Detailed Results' and 'MasterModels TP-1D' must both be found."""
    names = ["Raw Detailed Results", "Chart1", "MasterModels TP-1D"]
    assert du.find_sheet(names, ("detailedresults",), "wb") == "Raw Detailed Results"
    assert du.find_sheet(names, ("mastermodels",), "wb") == "MasterModels TP-1D"
    with pytest.raises(KeyError, match="Requirement 2"):
        du.find_sheet(["Sheet1"], ("detailedresults",), "wb")


def test_interpolation_is_linear_between_the_two_ends(workdir):
    """value = (1 - w) x ERR A mean + w x ERR B mean, at every one of the eleven steps."""
    du.build_self_test_workbook(workdir / "in")
    result = du.run_optimizer(data_dir=workdir / "in", out_dir=workdir / "out",
                              validate=False, verbose=False)
    computed = result["computed"]

    for base in (1, 1 + du.BLOCK_STRIDE):
        pos = du.block_rows(base)
        hdr = pos["dual_header"][du.SEX_ORDER[0]]
        row = hdr + 1                                   # first interpolation row of the block
        left = computed[(row, du.COL["H"])]
        right = computed[(row, du.COL["R"])]
        for i, w in enumerate(du.DUAL_USE_STEPS):
            assert computed[(row, du.COL["H"] + i)] == pytest.approx((1 - w) * left + w * right)


def test_tipping_point_bracket_is_found_not_assumed(workdir):
    """The crossing pair is located from the values, so new data does not need re-pointing."""
    du.build_self_test_workbook(workdir / "in")
    result = du.run_optimizer(data_dir=workdir / "in", out_dir=workdir / "out",
                              validate=False, verbose=False)
    tps = du.summarise_tipping_points(result)
    assert len(tps) == len(du.TP_SEXES) * len(result["loaded"]["gateways"]) * len(du.MATCHED_PAIRS)
    for value in tps["In combination with cigarettes"]:
        assert 0.0 <= float(value.rstrip("%")) <= 100.0
