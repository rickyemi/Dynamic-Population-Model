"""Workflow 1 and 2 on synthetic exports.

These call the module's own self test helpers, which build mock workbooks with the same tab names,
column names and model naming convention as the real DPM exports. Testing against generated inputs
rather than a committed fixture keeps product data out of the repository and still exercises the
parsing, which is where the naming conventions bite.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from dpm import relapse_popextrapolation as rp


def test_workflow1_self_test(workdir):
    rp.self_test_workflow1()


def test_workflow2_self_test(workdir):
    rp.self_test_workflow2()


def test_workflow1_writes_three_tabs(workdir):
    from openpyxl import load_workbook

    path = rp.w1_build_mock_source(workdir / "mock")
    result = rp.relapse_calculations(path.parent, path.name, "out.xlsx")
    tabs = load_workbook(result["output_file"]).sheetnames
    assert tabs == [rp.W1_SOURCE_SHEET, rp.W1_PIVOT_SHEET, rp.W1_DETAIL_SHEET]


def test_pivot_calc_columns_reconcile(workdir):
    """S = master ALL.T - relapse ALL.T, and T = master DIFF_ALL - S."""
    path = rp.w1_build_mock_source(workdir / "mock")
    result = rp.relapse_calculations(path.parent, path.name, "out.xlsx")
    piv = result["pivot_and_calcs"]

    expected_s = piv[f"{rp.W1_NODE_TOTAL}_m"] - piv[rp.W1_NODE_TOTAL]
    expected_t = piv[f"{rp.W1_NODE_DIFF}_m"] - expected_s
    assert np.allclose(piv["(CF1-CF2)"], expected_s)
    assert np.allclose(piv["CF1-BC1-CF1+CF2"], expected_t)


def test_counts_are_whole_numbers(workdir):
    """Survivor counts are whole people, rounded in the frame and the formula together."""
    path = rp.w1_build_mock_source(workdir / "mock")
    result = rp.relapse_calculations(path.parent, path.name, "out.xlsx")

    for col in ("(CF1-CF2)", "CF1-BC1-CF1+CF2", rp.W1_NODE_TOTAL):
        values = pd.to_numeric(result["pivot_and_calcs"][col]).dropna()
        assert (values % 1 == 0).all(), f"{col} still holds fractional values"
    for col in ("Mean_Webtool", "Mean"):
        values = pd.to_numeric(result["detailed_results"][col]).dropna()
        assert (values % 1 == 0).all(), f"{col} still holds fractional values"


def test_relapse_status_has_exactly_two_levels(workdir):
    """The no relapse arm carries no relapse wording, which used to raise."""
    frame = pd.DataFrame({"Group": ["Velo Max YO"] * 2,
                          "Name": ["Master Model G10 Velo Max Relapse_TwoAge",
                                   "Master Model G10 Velo Max YO_TwoAge_5a_6a_14b_15a"]})
    out = rp.classify_relapse(frame)
    assert list(out["Relapse.Status"].cat.categories) == rp.RELAPSE_LEVELS
    assert out["Relapse.Status"].tolist() == ["Relapse", "No Relapse"]


@pytest.mark.parametrize("group, product", [
    ("Grizzly XL YO", "Grizzly XL"), ("Velo_Max_YO", "Velo_Max"), ("Glo YO", "Glo"),
])
def test_product_extracted_from_model_group(group, product):
    found, used = rp.extract_product_name(pd.Series([group]))
    assert rp.normalise_product(found) == rp.normalise_product(product)
    assert used == group


def test_product_matching_ignores_spacing():
    assert rp.normalise_product("Grizzly XL") == rp.normalise_product("grizzly_xl")
    assert rp.normalise_product("Velo Max") != rp.normalise_product("Velo Min")
